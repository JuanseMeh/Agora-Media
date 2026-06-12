from __future__ import annotations

import logging
import os
import tempfile
from typing import Any

from config.settings import settings

logger = logging.getLogger(__name__)


async def parse_document_to_markdown(
    file_data: bytes,
    original_filename: str,
    mime_type: str,
) -> dict[str, Any]:
    ext = os.path.splitext(original_filename)[1].lower()

    try:
        if ext == ".pdf":
            return await _parse_pdf(file_data)
        elif ext in (".docx", ".doc"):
            return await _parse_docx(file_data)
        elif ext in (".pptx", ".ppt"):
            return await _parse_pptx(file_data)
        elif ext == ".txt":
            return await _parse_text(file_data)
        elif ext in (".md", ".markdown"):
            return {"markdown": file_data.decode("utf-8", errors="replace"), "method": "direct"}
        elif ext in (".html", ".htm"):
            return await _parse_html(file_data)
        elif ext in (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp"):
            return await _parse_image_text(file_data)
        elif mime_type.startswith("text/"):
            return await _parse_text(file_data)
        else:
            return {"error": f"Unsupported file type: {ext}", "markdown": ""}
    except Exception as e:
        logger.error("Document parsing failed: %s", e)
        return {"error": str(e), "markdown": ""}


async def _parse_pdf(data: bytes) -> dict[str, Any]:
    try:
        from marker.convert import convert_single_pdf
        from marker.models import load_all_models

        model_lst = load_all_models()

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name

        try:
            full_text, _, _ = convert_single_pdf(tmp_path, model_lst)
            markdown = full_text if full_text else ""
        finally:
            os.unlink(tmp_path)

        return {
            "markdown": markdown,
            "method": "marker",
            "char_count": len(markdown),
        }
    except ImportError:
        logger.warning("marker-pdf not available, falling back to pypdf")

    try:
        import pypdf

        reader = pypdf.PdfReader(tempfile.NamedTemporaryFile(suffix=".pdf"))
        text_parts = []
        for page in reader.pages:
            text_parts.append(page.extract_text())

        markdown = "\n\n".join(text_parts)
        return {
            "markdown": markdown,
            "method": "pypdf",
            "char_count": len(markdown),
        }
    except ImportError:
        return {"error": "No PDF parser available (install marker-pdf or pypdf)", "markdown": ""}


async def _parse_docx(data: bytes) -> dict[str, Any]:
    try:
        from docx import Document

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name

        try:
            doc = Document(tmp_path)
            markdown_parts = []

            for para in doc.paragraphs:
                style = para.style.name.lower() if para.style else ""
                text = para.text.strip()
                if not text:
                    continue

                if "heading 1" in style:
                    markdown_parts.append(f"# {text}")
                elif "heading 2" in style:
                    markdown_parts.append(f"## {text}")
                elif "heading 3" in style:
                    markdown_parts.append(f"### {text}")
                elif "list" in style:
                    markdown_parts.append(f"- {text}")
                else:
                    markdown_parts.append(text)

            for table in doc.tables:
                md_table = _table_to_markdown(table)
                if md_table:
                    markdown_parts.append(md_table)

            markdown = "\n\n".join(markdown_parts)
        finally:
            os.unlink(tmp_path)

        return {
            "markdown": markdown,
            "method": "python-docx",
            "char_count": len(markdown),
        }
    except ImportError:
        return {"error": "python-docx not installed", "markdown": ""}


async def _parse_pptx(data: bytes) -> dict[str, Any]:
    try:
        from pptx import Presentation

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name

        try:
            prs = Presentation(tmp_path)
            markdown_parts = []

            for i, slide in enumerate(prs.slides, 1):
                markdown_parts.append(f"## Slide {i}")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                markdown_parts.append(text)
                    if shape.has_table:
                        md_table = _table_to_markdown(shape.table)
                        if md_table:
                            markdown_parts.append(md_table)
                markdown_parts.append("")

            markdown = "\n\n".join(markdown_parts)
        finally:
            os.unlink(tmp_path)

        return {
            "markdown": markdown,
            "method": "python-pptx",
            "char_count": len(markdown),
        }
    except ImportError:
        return {"error": "python-pptx not installed", "markdown": ""}


async def _parse_text(data: bytes) -> dict[str, Any]:
    text = data.decode("utf-8", errors="replace")
    return {
        "markdown": text,
        "method": "text",
        "char_count": len(text),
    }


async def _parse_html(data: bytes) -> dict[str, Any]:
    try:
        import html2text

        html_content = data.decode("utf-8", errors="replace")
        converter = html2text.HTML2Text()
        converter.body_width = 0
        markdown = converter.handle(html_content)

        return {
            "markdown": markdown,
            "method": "html2text",
            "char_count": len(markdown),
        }
    except ImportError:
        text = data.decode("utf-8", errors="replace")
        return {
            "markdown": text,
            "method": "text_fallback",
            "char_count": len(text),
        }


async def _parse_image_text(data: bytes) -> dict[str, Any]:
    from services.vision import ocr_handwriting
    result = await ocr_handwriting(data)
    return {
        "markdown": result.get("text", ""),
        "method": "vision_ocr",
        "char_count": len(result.get("text", "")),
        "ocr_details": result,
    }


def _table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)

    if not rows:
        return ""

    md_lines = []
    md_lines.append("| " + " | ".join(rows[0]) + " |")
    md_lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
    for row in rows[1:]:
        md_lines.append("| " + " | ".join(row) + " |")

    return "\n".join(md_lines)
